"""Lot 70 — format « DECL » : la déclaration récapitulative des salaires au MODÈLE STANDARD de la CNSS.

Le modèle (feuille « DECL » du classeur du propriétaire) est conservé tel quel dans `modeles/decl_cnss.xlsx` : colonnes,
hauteurs de lignes, cellules fusionnées, bordures simples et doubles, polices, texte du rappel, MR / TG, cadre de signature.
Rien n'est redessiné « à la main » : le PDF est produit en LISANT ce gabarit cellule par cellule (aucune disposition
n'est écrite dans le code), puis en y inscrivant les valeurs de la déclaration. Le classeur Excel exporté part du même gabarit.

Logo : emplacement d'origine (en haut à gauche, colonne A, lignes 1 à 5). Le propriétaire peut y placer le logo des impôts
(téléversement dans l'onglet DECLARATION) ; sans logo téléversé, l'emblème de la CNSS du modèle d'origine est conservé.

Ce module est pur (aucune base de données) ; le stockage du logo et les routes sont dans albarka_paie_cnss.py.
"""
from __future__ import annotations

import io
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MODELES = Path(__file__).resolve().parent / "modeles"
GABARIT = MODELES / "decl_cnss.xlsx"               # feuille DECL d'origine (libellés et mise en forme)
LOGO_CNSS = MODELES / "logo_cnss.png"              # emblème CNSS du modèle d'origine (logo par défaut)
# Emplacement du logo (points, depuis le coin haut-gauche de la feuille) : colonne A, lignes 1 à 5 du modèle
CADRE_LOGO = (2.0, 1.5, 81.0, 73.0)                # x, y, largeur, hauteur

# Couleurs : remplissage « indexé 22 » du modèle (gris argent C0C0C0), ÉCLAIRCI à la demande du propriétaire (lot 71) en DBDBDB ;
# police « thème 10 » / « indexé 12 » (bleu des liens)
GRIS_CLAIR_HEX = "DBDBDB"
GRIS_ARGENT = (0xDB / 255, 0xDB / 255, 0xDB / 255)
BLEU = (0.0, 0.0, 1.0)


def logo_par_defaut() -> bytes:
    """Emblème de la CNSS du modèle d'origine (PNG)."""
    return LOGO_CNSS.read_bytes()


# --------------------------------------------------------------------------
# Valeurs inscrites dans le modèle
# --------------------------------------------------------------------------
def _iso_en_date(iso: str) -> date:
    return datetime.strptime(iso[:10], "%Y-%m-%d").date()


def valeurs_decl(d: Dict[str, Any]) -> Dict[str, Any]:
    """Cellules du modèle remplies automatiquement (mêmes cellules que les formules du classeur d'origine)."""
    e = d["employer"]
    N = d["travailleurs"]
    b = d["branches"]
    jrn_appr = N["journaliers"] + N["apprentis"]
    v: Dict[str, Any] = {
        "G8": _iso_en_date(d["debut"]), "J8": _iso_en_date(d["fin"]),            # période : DU … AU …
        "F10": e.get("name") or "",                                                  # société
        "D13": " — ".join(x for x in [e.get("address"), f"BP {e['bp']}" if e.get("bp") else ""] if x),
        "E14": f"TEL: {e['phone']}" if e.get("phone") else "",
        "I18": e.get("cnss_number") or "",                                           # matricule de l'employeur
        "B26": N["permanents"],
        "D26": N["temporaires"] or None, "F26": N["journaliers"] or None, "I26": N["apprentis"] or None,
        "H32": b[0]["taux"] / 100, "H35": b[1]["taux"] / 100, "H38": b[2]["taux"] / 100,
        "E31": b[0]["remuneration"], "I31": b[0]["cotisation"],
        "E34": b[1]["remuneration"], "I34": b[1]["cotisation"],
        "E37": b[2]["remuneration"], "I37": b[2]["cotisation"],
        "I41": d["cotisations"],
        "I51": _iso_en_date(d["date_certification"]),
        "A50": d["declarant"] or "",
    }
    for ligne in (32, 35, 38):                                                       # nombre de salariés par branche
        v[f"B{ligne}"], v[f"C{ligne}"], v[f"D{ligne}"] = N["permanents"], N["temporaires"], jrn_appr
    return v


# --------------------------------------------------------------------------
# Rendu PDF à partir du gabarit
# --------------------------------------------------------------------------
def _largeur_colonne_pt(ws, idx: int) -> float:
    """Largeur d'une colonne en points (Excel : pixels = largeur × 7, 1 pixel = 0,75 point)."""
    for dim in ws.column_dimensions.values():
        if dim.min and dim.max and dim.min <= idx <= dim.max and dim.width:
            return int(dim.width * 7 + 0.5) * 0.75
    return int(8.43 * 7 + 0.5) * 0.75


def _hauteur_ligne_pt(ws, idx: int) -> float:
    h = ws.row_dimensions[idx].height if idx in ws.row_dimensions else None
    return float(h or ws.sheet_format.defaultRowHeight or 12.75)


def _police(cell) -> Tuple[str, float, Tuple[float, float, float], bool]:
    """(nom de police PDF, taille, couleur, souligné) d'une cellule."""
    f = cell.font
    nom = (f.name or "Arial")
    gras, ital = bool(f.b), bool(f.i)
    if nom.lower().startswith("cambria"):
        pdf = "Times-BoldItalic" if gras and ital else "Times-Bold" if gras else "Times-Italic" if ital else "Times-Roman"
    else:
        gras = gras or nom.lower() == "arial black"
        pdf = "Helvetica-BoldOblique" if gras and ital else "Helvetica-Bold" if gras else "Helvetica-Oblique" if ital else "Helvetica"
    couleur = (0.0, 0.0, 0.0)
    c = f.color
    if c is not None and ((c.type == "theme" and c.theme == 10) or (c.type == "indexed" and c.indexed == 12)):
        couleur = BLEU
    return pdf, float(f.sz or 10), couleur, bool(f.u)


def _texte(valeur: Any, cell) -> Tuple[str, bool]:
    """(texte affiché, est un nombre) selon le format de la cellule (présentation française)."""
    if valeur is None:
        return "", False
    fmt = cell.number_format or "General"
    if isinstance(valeur, (datetime, date)):
        return (valeur.date() if isinstance(valeur, datetime) else valeur).strftime("%d/%m/%Y"), False
    if isinstance(valeur, bool):
        return str(valeur), False
    if isinstance(valeur, (int, float)):
        if "%" in fmt:
            n = 1 if "0.0" in fmt else 0
            return f"{valeur * 100:.{n}f}".replace(".", ",") + "%", True
        if "#,##0" in fmt:
            return ("-" if round(valeur) == 0 else f"{int(round(valeur)):,}".replace(",", " ")), True
        return (str(int(valeur)) if float(valeur).is_integer() else str(valeur).replace(".", ",")), True
    return str(valeur), False


def _trait(c, x1, y1, x2, y2, style: str, vertical: bool) -> None:
    """Trait de bordure : simple (0,5 pt) ou double (deux traits fins de part et d'autre)."""
    c.setStrokeColorRGB(0, 0, 0)
    if style == "double":
        c.setLineWidth(0.45)
        dx, dy = (0.7, 0) if vertical else (0, 0.7)
        c.line(x1 - dx, y1 - dy, x2 - dx, y2 - dy)
        c.line(x1 + dx, y1 + dy, x2 + dx, y2 + dy)
    else:
        c.setLineWidth(0.5 if style in ("thin", "hair") else 0.9)
        c.line(x1, y1, x2, y2)


def rendre_decl_pdf(valeurs: Dict[str, Any], logo: Optional[bytes] = None) -> bytes:
    """PDF A4 portrait de la feuille DECL du modèle, remplie avec `valeurs` (cellule → valeur) et le logo."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.pdfgen import canvas

    ws = load_workbook(GABARIT)["DECL"]
    nb_col = ws.max_column
    nb_lig = ws.max_row
    larg = [_largeur_colonne_pt(ws, i) for i in range(1, nb_col + 1)]
    haut = [_hauteur_ligne_pt(ws, i) for i in range(1, nb_lig + 1)]
    xs = [0.0]
    for w in larg:
        xs.append(xs[-1] + w)
    ys = [0.0]                                                   # ordonnées mesurées vers le BAS depuis le haut de la feuille
    for h in haut:
        ys.append(ys[-1] + h)
    total_l, total_h = xs[-1], ys[-1]

    # Page : marges du modèle ; réduction éventuelle pour tenir sur une page (« ajuster à une page » du modèle)
    page_l, page_h = A4
    m = ws.page_margins
    marge_g, marge_d, marge_h, marge_b = m.left * 72, m.right * 72, m.top * 72, m.bottom * 72
    echelle = min(1.0, (page_l - marge_g - marge_d) / total_l, (page_h - marge_h - marge_b) / total_h)

    tampon = io.BytesIO()
    c = canvas.Canvas(tampon, pagesize=A4)
    c.setTitle("Déclaration récapitulative des salaires (DECL)")
    c.translate(marge_g, page_h - marge_h)
    c.scale(echelle, echelle)
    c.translate(0, 0)

    def X(col: int) -> float:            # abscisse du bord gauche de la colonne (1 = A)
        return xs[col - 1]

    def Y(lig: int) -> float:            # ordonnée PDF du bord HAUT de la ligne (1 = première)
        return -ys[lig - 1]

    # Cellules fusionnées : coin haut-gauche → (lignes, colonnes) ; cellules masquées par une fusion
    fusion: Dict[Tuple[int, int], Tuple[int, int]] = {}
    masquees = set()
    for r in ws.merged_cells.ranges:
        fusion[(r.min_row, r.min_col)] = (r.max_row - r.min_row + 1, r.max_col - r.min_col + 1)
        for i in range(r.min_row, r.max_row + 1):
            for j in range(r.min_col, r.max_col + 1):
                if (i, j) != (r.min_row, r.min_col):
                    masquees.add((i, j))

    valeur_de = lambda cell: valeurs[cell.coordinate] if cell.coordinate in valeurs else cell.value  # noqa: E731

    # 1) remplissages (gris argent d'origine) ; dans une fusion, c'est celui de la cellule haut-gauche qui couvre toute la zone
    for ligne in ws.iter_rows(min_row=1, max_row=nb_lig, max_col=nb_col):
        for cell in ligne:
            if (cell.row, cell.column) in masquees:
                continue
            if cell.fill and cell.fill.fill_type == "solid":
                fg = cell.fill.fgColor
                if fg.type == "indexed" and fg.indexed == 22:
                    nl, nc = fusion.get((cell.row, cell.column), (1, 1))
                    c.setFillColorRGB(*GRIS_ARGENT)
                    c.rect(X(cell.column), Y(cell.row + nl - 1) - haut[cell.row + nl - 2],
                           X(cell.column + nc - 1) + larg[cell.column + nc - 2] - X(cell.column),
                           Y(cell.row) - (Y(cell.row + nl - 1) - haut[cell.row + nl - 2]), stroke=0, fill=1)

    # 2) bordures : chaque cellule ; dans une zone fusionnée, seuls les bords extérieurs comptent (comme dans Excel)
    zone_de: Dict[Tuple[int, int], Any] = {}
    for r in ws.merged_cells.ranges:
        for i in range(r.min_row, r.max_row + 1):
            for j in range(r.min_col, r.max_col + 1):
                zone_de[(i, j)] = r
    for ligne in ws.iter_rows(min_row=1, max_row=nb_lig, max_col=nb_col):
        for cell in ligne:
            z = zone_de.get((cell.row, cell.column))
            b = cell.border
            x1, x2 = X(cell.column), X(cell.column) + larg[cell.column - 1]
            yh, yb = Y(cell.row), Y(cell.row) - haut[cell.row - 1]
            if b.left and b.left.style and (z is None or cell.column == z.min_col):
                _trait(c, x1, yb, x1, yh, b.left.style, True)
            if b.right and b.right.style and (z is None or cell.column == z.max_col):
                _trait(c, x2, yb, x2, yh, b.right.style, True)
            if b.top and b.top.style and (z is None or cell.row == z.min_row):
                _trait(c, x1, yh, x2, yh, b.top.style, False)
            if b.bottom and b.bottom.style and (z is None or cell.row == z.max_row):
                _trait(c, x1, yb, x2, yb, b.bottom.style, False)

    # 3) textes et valeurs
    for ligne in ws.iter_rows(min_row=1, max_row=nb_lig, max_col=nb_col):
        for cell in ligne:
            if (cell.row, cell.column) in masquees:
                continue
            texte, est_nombre = _texte(valeur_de(cell), cell)
            if texte == "":
                continue
            nl, nc = fusion.get((cell.row, cell.column), (1, 1))
            x1 = X(cell.column)
            x2 = X(cell.column + nc - 1) + larg[cell.column + nc - 2]
            yh = Y(cell.row)
            yb = Y(cell.row + nl - 1) - haut[cell.row + nl - 2]
            nom, taille, couleur, souligne = _police(cell)
            c.setFont(nom, taille)
            c.setFillColorRGB(*couleur)
            # formats comptables : le chiffre est cadré à droite sauf alignement imposé
            horiz = cell.alignment.horizontal or ("right" if est_nombre else "left")
            if "#,##0" in (cell.number_format or "") and cell.alignment.horizontal is None:
                horiz = "right"
            marge = 2.0 + (stringWidth("F", nom, taille) if "#,##0" in (cell.number_format or "") and horiz == "right" else 0)
            w = stringWidth(texte, nom, taille)
            if horiz == "center":
                x = (x1 + x2) / 2 - w / 2
            elif horiz == "right":
                x = x2 - marge - w
            else:
                x = x1 + marge
            vert = cell.alignment.vertical or "bottom"
            if vert == "center":
                y = (yh + yb) / 2 - taille * 0.35
            elif vert == "top":
                y = yh - taille * 0.95
            else:
                y = yb + 2.2
            c.drawString(x, y, texte)
            if souligne:
                c.setStrokeColorRGB(*couleur)
                c.setLineWidth(0.5)
                c.line(x, y - 1.2, x + w, y - 1.2)

    # 4) logo à l'emplacement d'origine (proportions conservées, centré dans le cadre)
    img = logo if logo else logo_par_defaut()
    try:
        lecteur = ImageReader(io.BytesIO(img))
        iw, ih = lecteur.getSize()
        bx, by, bw, bh = CADRE_LOGO
        k = min(bw / iw, bh / ih)
        c.drawImage(lecteur, bx + (bw - iw * k) / 2, -(by + bh) + (bh - ih * k) / 2, iw * k, ih * k, mask="auto")
    except Exception:  # noqa: BLE001 — un logo illisible ne doit jamais empêcher l'impression
        pass

    c.showPage()
    c.save()
    return tampon.getvalue()


# --------------------------------------------------------------------------
# Excel : le gabarit rempli (formules d'origine conservées) + logo
# --------------------------------------------------------------------------
def remplir_feuille_decl_excel(ws, d: Dict[str, Any], ligne_total_btn: int, logo: Optional[bytes] = None) -> None:
    """Inscrit la déclaration dans la feuille DECL du gabarit, avec les formules du classeur d'origine."""
    from openpyxl.drawing.image import Image as XlImage

    from openpyxl.styles import PatternFill
    for ligne in ws.iter_rows():                                          # gris du rappel éclairci (comme sur le PDF)
        for cell in ligne:
            if cell.fill and cell.fill.fill_type == "solid" and cell.fill.fgColor.type == "indexed" and cell.fill.fgColor.indexed == 22:
                cell.fill = PatternFill("solid", fgColor=GRIS_CLAIR_HEX)
    v = valeurs_decl(d)
    for cle in ("G8", "J8", "F10", "D13", "E14", "I18", "B26", "D26", "F26", "I26", "H32", "H35", "H38", "A50", "I51"):
        ws[cle] = v[cle]
    ws["G8"].number_format = ws["J8"].number_format = ws["I51"].number_format = "dd/mm/yyyy"
    ref = f"=BTN!$D${ligne_total_btn}"                                   # total des rémunérations du BTN
    formules = {
        "E31": ref, "I31": "=ROUND(E31*H32,0)", "B32": "=B26", "C32": "=D26", "D32": "=F26+I26",
        "E34": "=SUM(E31)", "I34": "=ROUND(E34*H35,0)", "B35": "=SUM(B32)", "C35": "=SUM(C32)", "D35": "=SUM(D32)",
        "E37": "=SUM(E31)", "I37": "=ROUND(E37*H38,0)", "B38": "=SUM(B32)", "C38": "=SUM(C32)", "D38": "=SUM(D32)",
        "I41": "=SUM(I31:J39)",
    }
    for cle, f in formules.items():
        ws[cle] = f
    try:
        image = XlImage(io.BytesIO(logo if logo else logo_par_defaut()))
        image.width = image.height = 88                                  # pixels : tient dans la colonne A, lignes 1 à 5
        ws.add_image(image, "A1")
    except Exception:  # noqa: BLE001
        pass
    ws.page_setup.orientation, ws.page_setup.paperSize = "portrait", 9
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def logo_valide(octets: bytes) -> Tuple[bytes, str]:
    """Vérifie et normalise un logo téléversé : image PNG ou JPEG, ramenée à 600 px au plus, rendue en PNG."""
    from PIL import Image

    try:
        with Image.open(io.BytesIO(octets)) as im:
            if im.format not in ("PNG", "JPEG"):
                raise ValueError("format")
            im = im.convert("RGBA")
            im.thumbnail((600, 600))
            fond = Image.new("RGB", im.size, (255, 255, 255))      # transparence → fond blanc (papier)
            fond.paste(im, mask=im.split()[3])
            sortie = io.BytesIO()
            fond.save(sortie, "PNG", optimize=True)
            return sortie.getvalue(), "image/png"
    except Exception as exc:  # noqa: BLE001
        raise ValueError("Image illisible : choisissez un fichier PNG ou JPEG") from exc
