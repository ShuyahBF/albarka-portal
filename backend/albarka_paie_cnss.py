"""Lot 68 — Paie : DÉCLARATIONS CNSS (feuilles « BTN » et « DECLARATION » — « DECL » dans le classeur Excel du propriétaire).

Deux états, remplis automatiquement à partir des BULLETINS DE PAIE VALIDÉS de l'employeur choisi :
  - BTN   : Bordereau Nominatif des Travailleurs salariés (un salarié par ligne : n° CNSS, rémunération soumise à
            cotisations sur la période, période, observation), total, taux et cotisations ;
  - DECLARATION : Déclaration récapitulative des salaires, imprimée au format « DECL » = MODÈLE STANDARD DE LA CNSS, sans
            aucun changement de disposition (lot 70, voir albarka_paie_decl.py), avec le logo des impôts en haut à gauche
            (téléversé par le propriétaire ; à défaut, l'emblème CNSS du modèle). Mêmes chiffres que le BTN.

Règles (explications du propriétaire) :
  - taux global 21,5 % : assurance vieillesse 14 %, risques professionnels 1,5 %, prestations familiales 6 % ;
  - moins de 20 travailleurs : déclaration TRIMESTRIELLE ; à partir de 20 : déclaration MENSUELLE ;
  - rémunération soumise = base CNSS plafonnée de chaque bulletin (plafond mensuel 800 000 F, déjà appliqué par le
    moteur de paie), additionnée sur les mois de la période.

Sorties : écran (JSON), PDF imprimables (BTN en paysage, DECL en portrait, ou les deux ensemble) et Excel (feuilles DECL et BTN).
Lecture seule : aucune donnée de paie n'est modifiée ; seuls le nom du déclarant, la ville et le logo sont mémorisés.
"""
from __future__ import annotations

import calendar
import io
import re
from datetime import date
from typing import Any, Dict, List, Optional
from xml.sax.saxutils import escape

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field

from albarka_auth import require_roles
from albarka_paie import CABINET_ID, HR_ROLES, _employer_identity, employer_config
from albarka_docgen import MONTHS_FR
from db import db

router = APIRouter(prefix="/hr/paie/cnss", tags=["RH & Paie — déclarations CNSS"])

# Branches de cotisation et leurs taux (feuille DECL du classeur : 14 % + 1,5 % + 6 % = 21,5 %)
BRANCHES = [("Assurance vieillesse", 14.0), ("Risques professionnels", 1.5), ("Prestations familiales", 6.0)]
TAUX_TOTAL = round(sum(t for _, t in BRANCHES), 2)          # 21,5
SEUIL_MENSUEL = 20                                          # à partir de 20 travailleurs : déclaration mensuelle
VILLE_DEFAUT = "OUAGADOUGOU"

# --------------------------------------------------------------------------
# Période
# --------------------------------------------------------------------------
def _mois_de_la_periode(annee: int, periodicite: str, numero: int) -> List[str]:
    """Mois (AAAA-MM) couverts : un mois (mensuelle) ou trois mois (trimestrielle)."""
    if periodicite == "mensuelle":
        if not 1 <= numero <= 12:
            raise HTTPException(status_code=422, detail="Mois attendu entre 1 et 12")
        return [f"{annee}-{numero:02d}"]
    if not 1 <= numero <= 4:
        raise HTTPException(status_code=422, detail="Trimestre attendu entre 1 et 4")
    premier = 3 * (numero - 1) + 1
    return [f"{annee}-{m:02d}" for m in range(premier, premier + 3)]


def _debut_fin(mois: List[str]) -> tuple[str, str]:
    """Premier et dernier jour de la période (AAAA-MM-JJ)."""
    a, m = map(int, mois[-1].split("-"))
    return f"{mois[0]}-01", f"{mois[-1]}-{calendar.monthrange(a, m)[1]:02d}"


def _fr(iso: Optional[str]) -> str:
    """AAAA-MM-JJ → JJ/MM/AAAA."""
    if not iso or len(iso) < 10:
        return ""
    return f"{iso[8:10]}/{iso[5:7]}/{iso[:4]}"


def _nombre(v: Any) -> str:
    """Montant en francs avec espace de milliers (« 1 234 567 »)."""
    return f"{int(round(float(v or 0))):,}".replace(",", " ")


def _taux(t: float) -> str:
    """Taux en % à la française, sans zéro inutile (14 → « 14 », 1.5 → « 1,5 »)."""
    return f"{t:g}".replace(".", ",")


def _libelle_periode(periodicite: str, annee: int, numero: int) -> str:
    return f"{MONTHS_FR[numero - 1]} {annee}" if periodicite == "mensuelle" else f"{numero}{'er' if numero == 1 else 'e'} trimestre {annee}"


# --------------------------------------------------------------------------
# Calcul de la déclaration
# --------------------------------------------------------------------------
async def construire_declaration(employer_id: str, annee: int, periodicite: str = "auto", numero: int = 1,
                                 inclure_brouillons: bool = False) -> Dict[str, Any]:
    """Données des deux états (BTN et DECLARATION) pour un employeur et une période."""
    if periodicite not in ("auto", "mensuelle", "trimestrielle"):
        raise HTTPException(status_code=422, detail="Périodicité inconnue")
    if not 2000 <= annee <= 2100:
        raise HTTPException(status_code=422, detail="Année invalide")
    cfg = await employer_config(employer_id, strict=False)
    # Effectif = salariés enregistrés de l'employeur : moins de 20 → trimestrielle, 20 et plus → mensuelle
    effectif = await db.employees.count_documents({"tenant_id": employer_id})
    auto = "mensuelle" if effectif >= SEUIL_MENSUEL else "trimestrielle"
    mode = auto if periodicite == "auto" else periodicite
    mois = _mois_de_la_periode(annee, mode, numero)
    debut, fin = _debut_fin(mois)

    # Bulletins de la période : validés (et, si demandé, brouillons)
    bulletins = await db.paie_bulletins.find({"employer_id": employer_id, "period_month": {"$in": mois}},
                                             {"_id": 0, "params": 0}).to_list(5000)
    retenus = [b for b in bulletins if inclure_brouillons or b.get("status") == "validated"]
    brouillons = sum(1 for b in bulletins if b.get("status") != "validated")

    # Regroupement par salarié : somme des bases CNSS plafonnées, mois concernés, dernière fiche connue
    par_salarie: Dict[str, Dict[str, Any]] = {}
    for b in sorted(retenus, key=lambda x: x["period_month"]):
        r = b.get("result") or {}
        base = float(r.get("employer_cnss_base") if r.get("employer_cnss_base") is not None else (r.get("cnss_base") or 0))
        s = par_salarie.setdefault(b["employee_id"], {"base": 0.0, "mois": [], "employee": {}})
        s["base"] += base
        s["mois"].append(b["period_month"])
        s["employee"] = b.get("employee") or s["employee"]
    lignes = []
    for s in sorted(par_salarie.values(), key=lambda x: (x["employee"].get("full_name") or "").lower()):
        e = s["employee"]
        d, f = _debut_fin(s["mois"])
        n = len(s["mois"])
        lignes.append({"full_name": e.get("full_name") or "", "cnss_number": (e.get("cnss_number") or "").strip() or "EN COURS",
                       "remuneration": int(round(s["base"])), "du": d, "au": f, "observation": f"{n} Mois"})
    for i, l in enumerate(lignes, start=1):
        l["n"] = i

    total = sum(l["remuneration"] for l in lignes)
    branches = [{"libelle": lib, "taux": taux, "remuneration": total, "cotisation": int(round(total * taux / 100))}
                for lib, taux in BRANCHES]
    total_cotisations = sum(b["cotisation"] for b in branches)

    # Alertes affichées à l'écran (rien n'est bloquant)
    alertes: List[str] = []
    if not cfg.get("cnss_number"):
        alertes.append("N° CNSS de l'employeur non renseigné : à saisir dans « Paramètres de paie ».")
    if not retenus:
        alertes.append("Aucun bulletin validé sur cette période : validez les bulletins du mois dans l'onglet « Bulletins de paie ».")
    if brouillons and not inclure_brouillons:
        alertes.append(f"{brouillons} bulletin(s) en brouillon ne sont pas comptés : validez-les ou cochez « Inclure les brouillons ».")
    employes = await db.employees.find({"tenant_id": employer_id}, {"_id": 0, "id": 1, "full_name": 1}).to_list(2000)
    sans_bulletin = [e.get("full_name") or "" for e in employes if e["id"] not in par_salarie]
    if retenus and sans_bulletin:
        alertes.append("Sans bulletin sur la période : " + ", ".join(sans_bulletin[:8]) + ("…" if len(sans_bulletin) > 8 else "") + ".")
    if any(l["cnss_number"] == "EN COURS" for l in lignes):
        alertes.append("Un ou plusieurs salariés n'ont pas de n° CNSS : « EN COURS » est imprimé à leur place.")

    reglages = await db.paie_employeurs.find_one({"employer_id": employer_id}, {"_id": 0, "cnss_declarant": 1, "cnss_ville": 1}) or {}
    return {
        "employer": {k: cfg.get(k) or "" for k in ("employer_id", "name", "address", "bp", "phone", "cnss_number")},
        "periodicite": mode, "periodicite_auto": auto, "effectif": effectif, "seuil": SEUIL_MENSUEL,
        "annee": annee, "numero": numero, "mois": mois, "debut": debut, "fin": fin,
        "libelle_periode": _libelle_periode(mode, annee, numero),
        "lignes": lignes, "total_remuneration": total, "taux_total": TAUX_TOTAL, "cotisations": total_cotisations,
        "branches": branches,
        "travailleurs": {"permanents": len(lignes), "temporaires": 0, "journaliers": 0, "apprentis": 0},
        "declarant": reglages.get("cnss_declarant") or "", "ville": reglages.get("cnss_ville") or VILLE_DEFAUT,
        "inclure_brouillons": inclure_brouillons, "alertes": alertes,
        "logo_personnalise": bool(await db.settings.find_one({"_id": LOGO_DOC_ID}, {"_id": 1})),
        "date_certification": date.today().isoformat(),
    }


async def _donnees(employer_id: str, annee: int, periodicite: str, numero: int, inclure_brouillons: bool) -> Dict[str, Any]:
    return await construire_declaration(employer_id, annee, periodicite, numero, inclure_brouillons)


@router.get("/declaration")
async def declaration(employer_id: str, annee: int, numero: int, periodicite: str = "auto", inclure_brouillons: bool = False,
                      user: dict = Depends(require_roles(HR_ROLES))):
    """Données de la déclaration CNSS (BTN et DECLARATION) pour l'écran."""
    return await _donnees(employer_id, annee, periodicite, numero, inclure_brouillons)


class DeclarantIn(BaseModel):
    nom: str = Field("", max_length=120)          # « Nom et prénoms » du déclarant (bas de la DECLARATION)
    ville: str = Field(VILLE_DEFAUT, max_length=80)   # « à OUAGADOUGOU » (bas du BTN)


@router.put("/{eid}/declarant")
async def enregistrer_declarant(eid: str, p: DeclarantIn, user: dict = Depends(require_roles(HR_ROLES))):
    """Mémorise le nom du déclarant et la ville imprimés en bas des états (par employeur)."""
    if eid != CABINET_ID:
        await _employer_identity(eid)             # 404 si l'employeur n'existe pas
    await db.paie_employeurs.update_one(
        {"employer_id": eid},
        {"$set": {"cnss_declarant": p.nom.strip(), "cnss_ville": (p.ville or "").strip() or VILLE_DEFAUT}},
        upsert=True)
    return {"ok": True}


# --------------------------------------------------------------------------
# Logo imprimé en haut à gauche de la DECL (logo des impôts téléversé par le propriétaire)
# --------------------------------------------------------------------------
LOGO_DOC_ID = "paie_cnss_logo"                  # document de la collection `settings` (une seule image pour tout le portail)
LOGO_MAX_OCTETS = 3 * 1024 * 1024               # taille maximale du fichier choisi (3 Mo) ; il est ensuite réduit à 600 px


async def _logo_octets() -> Optional[bytes]:
    """Logo téléversé (PNG), ou None : l'emblème CNSS du modèle d'origine est alors utilisé."""
    import base64
    doc = await db.settings.find_one({"_id": LOGO_DOC_ID}, {"_id": 0, "data": 1})
    return base64.b64decode(doc["data"]) if doc and doc.get("data") else None


@router.get("/logo")
async def lire_logo(user: dict = Depends(require_roles(HR_ROLES))):
    """Logo actuellement imprimé (téléversé, sinon emblème CNSS du modèle). En-tête X-Logo : personnalise | modele."""
    from albarka_paie_decl import logo_par_defaut
    octets = await _logo_octets()
    return Response(octets or logo_par_defaut(), media_type="image/png", headers={"X-Logo": "personnalise" if octets else "modele",
                                                                                 "Cache-Control": "no-store"})


@router.post("/logo")
async def televerser_logo(fichier: UploadFile = File(...), user: dict = Depends(require_roles(HR_ROLES))):
    """Enregistre le logo des impôts (PNG ou JPEG) à imprimer en haut à gauche de la DECL."""
    import base64

    from albarka_paie_decl import logo_valide
    brut = await fichier.read(LOGO_MAX_OCTETS + 1)
    if len(brut) > LOGO_MAX_OCTETS:
        raise HTTPException(status_code=413, detail="Image trop lourde (3 Mo au plus)")
    try:
        png, _ = logo_valide(brut)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc))
    await db.settings.update_one({"_id": LOGO_DOC_ID}, {"$set": {"data": base64.b64encode(png).decode("ascii"),
                                 "nom": (fichier.filename or "")[:120], "updated_at": date.today().isoformat(),
                                 "updated_by": user.get("email") or user.get("id")}}, upsert=True)
    return {"ok": True}


@router.delete("/logo")
async def retablir_logo(user: dict = Depends(require_roles(HR_ROLES))):
    """Revient à l'emblème CNSS du modèle d'origine."""
    await db.settings.delete_one({"_id": LOGO_DOC_ID})
    return {"ok": True}


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def _styles():
    from reportlab.lib.styles import ParagraphStyle
    return {
        "p": ParagraphStyle("p", fontName="Helvetica", fontSize=9, leading=11),
        "pb": ParagraphStyle("pb", fontName="Helvetica-Bold", fontSize=9, leading=11),
        "petit": ParagraphStyle("petit", fontName="Helvetica", fontSize=7.8, leading=9.6),
        "titre": ParagraphStyle("titre", fontName="Helvetica-Bold", fontSize=13, leading=16, alignment=1),
        "centre": ParagraphStyle("centre", fontName="Helvetica", fontSize=9, leading=11, alignment=1),
        "centreb": ParagraphStyle("centreb", fontName="Helvetica-Bold", fontSize=10, leading=12, alignment=1),
    }


def build_btn_pdf(d: Dict[str, Any]) -> bytes:
    """BTN : Bordereau Nominatif des Travailleurs salariés (A4 paysage, plusieurs pages si besoin)."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    st = _styles()
    e = d["employer"]
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=10 * mm,
                            bottomMargin=10 * mm, title=f"BTN CNSS {e['name']} {d['libelle_periode']}")
    gris = colors.HexColor("#CBD5E1")

    # En-tête : période à gauche, titre + société au milieu, n° CNSS à droite
    gauche = [Paragraph("<b>Période de déclaration</b>", st["p"]), Paragraph(f"du {_fr(d['debut'])}", st["p"]),
              Paragraph(f"au {_fr(d['fin'])}", st["p"])]
    milieu = [Paragraph("Bordereau Nominatif des Travailleurs Salariés", st["titre"]), Spacer(1, 3),
              Paragraph(escape((e.get("name") or "").upper()), st["centreb"])]
    if e.get("address") or e.get("bp"):
        milieu.append(Paragraph(escape(" — ".join(x for x in [e.get("address"), f"BP {e['bp']}" if e.get("bp") else ""] if x)), st["centre"]))
    if e.get("phone"):
        milieu.append(Paragraph(escape(f"TEL: {e['phone']}"), st["centre"]))
    droite = [Paragraph("<b>NUMERO :</b>", st["p"]), Paragraph(escape(e.get("cnss_number") or "…………"), st["pb"])]
    entete = Table([[gauche, milieu, droite]], colWidths=[62 * mm, 130 * mm, 65 * mm])
    entete.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOX", (0, 0), (-1, -1), 0.6, colors.black),
                                ("LINEAFTER", (0, 0), (1, 0), 0.6, colors.black), ("ALIGN", (2, 0), (2, 0), "RIGHT")]))

    # Tableau nominatif
    tete = ["N°\nOrd.", "NOM & PRÉNOMS", "N° D'IMMATRICULATION\nCNSS", "RÉMUNÉRATIONS\nSoumises à cotisations", "PÉRIODE\nDu … au …", "OBSERVATIONS"]
    rows: List[List[Any]] = [tete]
    for l in d["lignes"]:
        rows.append([str(l["n"]), Paragraph(escape(l["full_name"]), st["p"]), l["cnss_number"], _nombre(l["remuneration"]),
                     f"{_fr(l['du'])} au {_fr(l['au'])}", l["observation"]])
    if not d["lignes"]:
        rows.append(["", "Aucun salarié sur cette période", "", "", "", ""])
    n = len(rows)
    rows += [["", "", "TOTAL :", _nombre(d["total_remuneration"]), "", ""],
             ["", "", "TAUX :", f"{_taux(d['taux_total'])} %", "", ""],
             ["", "", "COTISATIONS :", _nombre(d["cotisations"]), "", ""]]
    t = Table(rows, colWidths=[14 * mm, 78 * mm, 48 * mm, 44 * mm, 47 * mm, 26 * mm], repeatRows=1)
    t.setStyle(TableStyle([
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8), ("ALIGN", (0, 0), (-1, 0), "CENTER"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F5F9")),
        ("FONT", (0, 1), (-1, -1), "Helvetica", 9), ("GRID", (0, 0), (-1, n - 1), 0.5, colors.black),
        ("ALIGN", (0, 1), (0, n - 1), "CENTER"), ("ALIGN", (3, 1), (3, -1), "RIGHT"), ("ALIGN", (2, 1), (2, n - 1), "CENTER"),
        ("ALIGN", (4, 1), (5, n - 1), "CENTER"),
        # trois lignes de bas de tableau : total, taux, cotisations
        ("FONT", (2, n), (3, -1), "Helvetica-Bold", 9.5), ("ALIGN", (2, n), (2, -1), "RIGHT"),
        ("BOX", (2, n), (3, -1), 0.8, colors.black), ("INNERGRID", (2, n), (3, -1), 0.5, gris),
        ("BACKGROUND", (2, n), (3, -1), colors.HexColor("#F8FAFC")),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    pied = Table([[Paragraph(f"Certifié exact le <b>{_fr(d['date_certification'])}</b>", st["p"]), ""],
                  [Paragraph(f"à {escape(d['ville'])}", st["p"]), Paragraph(escape(d["declarant"] or ""), st["pb"])]],
                 colWidths=[130 * mm, 127 * mm])
    pied.setStyle(TableStyle([("ALIGN", (1, 0), (1, -1), "CENTER"), ("TOPPADDING", (0, 0), (-1, -1), 2)]))
    doc.build([entete, Spacer(1, 6), t, Spacer(1, 10), pied])
    return buf.getvalue()


def build_recap_pdf(d: Dict[str, Any], logo: Optional[bytes] = None) -> bytes:
    """DECLARATION au format « DECL » : modèle standard de la CNSS, disposition inchangée (voir albarka_paie_decl.py)."""
    from albarka_paie_decl import rendre_decl_pdf, valeurs_decl
    return rendre_decl_pdf(valeurs_decl(d), logo)


def _nom_fichier(prefixe: str, d: Dict[str, Any], ext: str) -> str:
    base = re.sub(r"[^A-Za-z0-9_-]+", "_", f"{prefixe}_{d['employer'].get('name') or 'employeur'}_{d['annee']}_{d['periodicite'][0].upper()}{d['numero']}")
    return f"{base[:90]}.{ext}"


@router.get("/btn/pdf")
async def btn_pdf(employer_id: str, annee: int, numero: int, periodicite: str = "auto", inclure_brouillons: bool = False,
                  user: dict = Depends(require_roles(HR_ROLES))):
    """BTN imprimable (PDF, A4 paysage)."""
    import asyncio
    d = await _donnees(employer_id, annee, periodicite, numero, inclure_brouillons)
    pdf = await asyncio.to_thread(build_btn_pdf, d)
    return Response(pdf, media_type="application/pdf", headers={"Content-Disposition": f'inline; filename="{_nom_fichier("BTN", d, "pdf")}"'})


@router.get("/recap/pdf")
async def recap_pdf(employer_id: str, annee: int, numero: int, periodicite: str = "auto", inclure_brouillons: bool = False,
                    user: dict = Depends(require_roles(HR_ROLES))):
    """DECLARATION imprimable (PDF, A4 portrait)."""
    import asyncio
    d = await _donnees(employer_id, annee, periodicite, numero, inclure_brouillons)
    pdf = await asyncio.to_thread(build_recap_pdf, d, await _logo_octets())
    return Response(pdf, media_type="application/pdf", headers={"Content-Disposition": f'inline; filename="{_nom_fichier("DECLARATION", d, "pdf")}"'})


@router.get("/pdf")
async def declaration_pdf(employer_id: str, annee: int, numero: int, periodicite: str = "auto", inclure_brouillons: bool = False,
                          user: dict = Depends(require_roles(HR_ROLES))):
    """Déclaration complète : BTN puis DECLARATION dans un seul PDF (impression en une fois)."""
    import asyncio

    import fitz  # PyMuPDF, déjà utilisé par le portail
    d = await _donnees(employer_id, annee, periodicite, numero, inclure_brouillons)
    logo = await _logo_octets()

    def fusion() -> bytes:
        out = fitz.open()
        for pdf in (build_btn_pdf(d), build_recap_pdf(d, logo)):
            with fitz.open(stream=pdf, filetype="pdf") as un:
                out.insert_pdf(un)
        return out.tobytes()
    pdf = await asyncio.to_thread(fusion)
    return Response(pdf, media_type="application/pdf", headers={"Content-Disposition": f'inline; filename="{_nom_fichier("CNSS", d, "pdf")}"'})


# --------------------------------------------------------------------------
# Excel : mêmes deux feuilles que le classeur du propriétaire (totaux en formules)
# --------------------------------------------------------------------------
def build_excel(d: Dict[str, Any], logo: Optional[bytes] = None) -> bytes:
    """Classeur à deux feuilles comme celui du propriétaire : DECL (gabarit d'origine, formules conservées) puis BTN."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    from albarka_paie_decl import GABARIT, remplir_feuille_decl_excel

    e = d["employer"]
    gras = Font(bold=True)
    fin = Side(style="thin")
    cadre = Border(left=fin, right=fin, top=fin, bottom=fin)
    entete = PatternFill("solid", fgColor="F1F5F9")
    wb = load_workbook(GABARIT)                          # feuille « DECL » : mise en forme et libellés du modèle CNSS

    # ---- Feuille BTN (paysage, ajustée à la largeur d'une page) ----
    ws = wb.create_sheet("BTN")
    ws["A1"], ws["A2"], ws["A3"] = "Période de déclaration", f"du {_fr(d['debut'])}", f"au {_fr(d['fin'])}"
    ws["C1"], ws["C2"] = "Bordereau Nominatif des Travailleurs Salariés", (e.get("name") or "").upper()
    ws["C3"] = " — ".join(x for x in [e.get("address"), f"BP {e['bp']}" if e.get("bp") else ""] if x)
    ws["C4"] = f"TEL: {e['phone']}" if e.get("phone") else ""
    ws["F3"], ws["G3"] = "NUMERO :", e.get("cnss_number") or ""
    for c in ("A1", "C1", "C2", "G3"):
        ws[c].font = gras
    for j, h in enumerate(["N° Ord.", "NOM & PRÉNOMS", "N° D'IMMATRICULATION CNSS", "RÉMUNÉRATIONS soumises à cotisations",
                           "PÉRIODE du", "au", "OBSERVATIONS"], start=1):
        c = ws.cell(row=6, column=j, value=h)
        c.font, c.fill, c.border = gras, entete, cadre
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ligne = 7
    for l in d["lignes"]:
        valeurs = [l["n"], l["full_name"], l["cnss_number"], l["remuneration"], _fr(l["du"]), _fr(l["au"]), l["observation"]]
        for j, v in enumerate(valeurs, start=1):
            c = ws.cell(row=ligne, column=j, value=v)
            c.border = cadre
            if j == 4:
                c.number_format = "#,##0"
        ligne += 1
    premiere, derniere = 7, max(7, ligne - 1)
    ws.cell(row=ligne, column=3, value="TOTAL :").font = gras
    ws.cell(row=ligne, column=4, value=f"=SUM(D{premiere}:D{derniere})").number_format = "#,##0"
    ws.cell(row=ligne + 1, column=3, value="TAUX :").font = gras
    ws.cell(row=ligne + 1, column=4, value=d["taux_total"] / 100).number_format = "0.0%"
    ws.cell(row=ligne + 2, column=3, value="COTISATIONS :").font = gras
    ws.cell(row=ligne + 2, column=4, value=f"=ROUND(D{ligne}*D{ligne + 1},0)").number_format = "#,##0"
    ws.cell(row=ligne + 4, column=1, value=f"Certifié exact le {_fr(d['date_certification'])}")
    ws.cell(row=ligne + 5, column=1, value=f"à {d['ville']}")
    ws.cell(row=ligne + 5, column=4, value=d["declarant"])
    for col, larg in zip("ABCDEFG", (9, 36, 26, 24, 13, 13, 16)):
        ws.column_dimensions[col].width = larg
    ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", 9
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "6:6"

    # ---- Feuille DECL : gabarit rempli (le total des rémunérations vient du BTN ci-dessus) ----
    remplir_feuille_decl_excel(wb["DECL"], d, ligne, logo)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/excel")
async def declaration_excel(employer_id: str, annee: int, numero: int, periodicite: str = "auto", inclure_brouillons: bool = False,
                            user: dict = Depends(require_roles(HR_ROLES))):
    """Classeur Excel à deux feuilles (BTN et DECLARATION), totaux calculés par formules."""
    import asyncio
    d = await _donnees(employer_id, annee, periodicite, numero, inclure_brouillons)
    data = await asyncio.to_thread(build_excel, d, await _logo_octets())
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{_nom_fichier("CNSS", d, "xlsx")}"'})
